"""
Ứng dụng web nội bộ tra cứu hóa đơn đầu vào / đầu ra từ Tổng cục Thuế.

Chạy:
    pip install -r requirements.txt
    uvicorn backend.main:app --reload --port 8000
Rồi mở http://localhost:8000

Không lưu mật khẩu ở server. Sau khi đăng nhập, chỉ có JWT token (do Tổng cục
Thuế cấp) được giữ trong bộ nhớ phiên để gọi API. Token này tự hết hạn.
"""

from __future__ import annotations

import io
import logging
import uuid
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from .gdt_client import (
    EXPORT_COLUMNS,
    GdtClient,
    GdtContractDriftError,
    GdtError,
    normalize_row,
)

logger = logging.getLogger("vatcrawlbot.main")

app = FastAPI(title="Tra cứu hóa đơn TCT - Nội bộ")

# Lưu phiên trong RAM: session_id -> {client, captcha_key}
# (Ứng dụng chạy nội bộ 1 máy; nếu triển khai nhiều tiến trình cần dùng Redis.)
_SESSIONS: Dict[str, Dict[str, Any]] = {}


def _get_client(session_id: Optional[str]) -> GdtClient:
    sess = _SESSIONS.get(session_id or "")
    if not sess or not sess.get("client") or not sess["client"].token:
        raise HTTPException(status_code=401, detail="Chưa đăng nhập hoặc phiên đã hết hạn.")
    return sess["client"]


# --------------------------- Models --------------------------- #
class LoginBody(BaseModel):
    session_id: str
    username: str
    password: str
    ckey: str
    cvalue: str
    verify_ssl: bool = True


class QueryBody(BaseModel):
    session_id: str
    direction: str  # purchase | sold
    date_from: str  # dd/mm/yyyy
    date_to: str  # dd/mm/yyyy
    include_sco: bool = True
    statuses: Optional[List[int]] = [5, 6, 8]


# --------------------------- API ------------------------------ #
@app.get("/api/captcha")
def api_captcha():
    """Tạo phiên mới + trả captcha."""
    session_id = uuid.uuid4().hex
    client = GdtClient()
    _SESSIONS[session_id] = {"client": client, "captcha_key": None}
    try:
        cap = client.get_captcha()
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Không lấy được captcha: {e}")
    _SESSIONS[session_id]["captcha_key"] = cap.get("key")
    return {"session_id": session_id, "key": cap.get("key"), "content": cap.get("content")}


@app.get("/api/captcha/refresh")
def api_captcha_refresh(session_id: str):
    sess = _SESSIONS.get(session_id)
    if not sess:
        raise HTTPException(status_code=400, detail="Phiên không tồn tại.")
    try:
        cap = sess["client"].get_captcha()
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Không lấy được captcha: {e}")
    sess["captcha_key"] = cap.get("key")
    return {"key": cap.get("key"), "content": cap.get("content")}


@app.post("/api/login")
def api_login(body: LoginBody):
    sess = _SESSIONS.get(body.session_id)
    if not sess:
        raise HTTPException(status_code=400, detail="Phiên không tồn tại. Hãy tải lại captcha.")
    client: GdtClient = sess["client"]
    client.verify_ssl = body.verify_ssl
    try:
        client.login(body.username, body.password, body.ckey, body.cvalue)
    except GdtContractDriftError as e:
        # Không phải "sai mật khẩu"/"hết phiên" — API Tổng cục Thuế có thể đã
        # đổi cấu trúc. Không trả 401 (sẽ khiến người dùng tưởng do họ nhập
        # sai và thử lại vô ích) — báo lỗi hệ thống rõ ràng.
        logger.critical("GDT contract drift khi đăng nhập: %s", e)
        raise HTTPException(status_code=503, detail=f"Hệ thống thuế có thay đổi bất thường: {e}")
    except GdtError as e:
        raise HTTPException(status_code=401, detail=str(e))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Lỗi kết nối máy chủ thuế: {e}")
    return {"ok": True, "message": "Đăng nhập thành công."}


@app.post("/api/invoices")
def api_invoices(body: QueryBody):
    client = _get_client(body.session_id)
    try:
        rows = client.query_invoices(
            direction=body.direction,
            date_from=body.date_from,
            date_to=body.date_to,
            statuses=body.statuses,
            include_sco=body.include_sco,
        )
    except GdtContractDriftError as e:
        logger.critical("GDT contract drift khi truy vấn hóa đơn: %s", e)
        raise HTTPException(status_code=503, detail=f"Hệ thống thuế có thay đổi bất thường: {e}")
    except GdtError as e:
        raise HTTPException(status_code=401, detail=str(e))
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Lỗi truy vấn: {e}")

    norm = [normalize_row(r) for r in rows]
    # Tổng hợp nhanh
    total_pretax = sum(float(r.get("tgtcthue") or 0) for r in rows)
    total_tax = sum(float(r.get("tgtthue") or 0) for r in rows)
    total_pay = sum(float(r.get("tgtttbso") or 0) for r in rows)
    return {
        "count": len(norm),
        "columns": [{"key": k, "label": lbl} for k, lbl in EXPORT_COLUMNS],
        "rows": norm,
        "summary": {
            "tong_chua_thue": total_pretax,
            "tong_thue": total_tax,
            "tong_thanh_toan": total_pay,
        },
    }


@app.post("/api/export")
def api_export(body: QueryBody):
    client = _get_client(body.session_id)
    rows = client.query_invoices(
        direction=body.direction,
        date_from=body.date_from,
        date_to=body.date_to,
        statuses=body.statuses,
        include_sco=body.include_sco,
    )
    norm = [normalize_row(r) for r in rows]

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "HoaDon"

    headers = [lbl for _, lbl in EXPORT_COLUMNS]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    money_keys = {"tgtcthue", "tgtthue", "tgtttbso"}
    for r in norm:
        ws.append([r.get(k, "") for k, _ in EXPORT_COLUMNS])

    # Định dạng cột tiền
    for idx, (k, _lbl) in enumerate(EXPORT_COLUMNS, start=1):
        col = ws.cell(row=1, column=idx).column_letter
        width = 14
        if k in ("nbten", "nmten"):
            width = 34
        elif k in ("tgtttbso", "tgtcthue", "tgtthue"):
            width = 16
        ws.column_dimensions[col].width = width
        if k in money_keys:
            for row_cells in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                for c in row_cells:
                    c.number_format = "#,##0"
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"hoadon_{body.direction}_{body.date_from.replace('/', '')}_{body.date_to.replace('/', '')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@app.post("/api/detail")
def api_detail(body: Dict[str, Any]):
    client = _get_client(body.get("session_id"))
    row = body.get("row") or {}
    try:
        return client.invoice_detail(row)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"Lỗi lấy chi tiết: {e}")


# --------------------------- Static (frontend) ---------------- #
import os

_FRONTEND_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "frontend")


@app.get("/")
def index():
    return FileResponse(os.path.join(_FRONTEND_DIR, "index.html"))


app.mount("/static", StaticFiles(directory=_FRONTEND_DIR), name="static")
